from pathlib import Path

from docx import Document

from talend_doc_gen.models import Component, Connection, Job
from talend_doc_gen.render_docx import render_job_docx


def make_job() -> Job:
    return Job(
        name="Demo_Job",
        description="Un job de démonstration.",
        owner="Equipe Data",
        components=[
            Component(name="tFileInputDelimited_1", type="tFileInputDelimited"),
            Component(name="tDBOutput_1", type="tDBOutput"),
        ],
        connections=[
            Connection(source="tFileInputDelimited_1", target="tDBOutput_1", label="Main", kind="flow"),
        ],
        context_params={"connection_sftp_password": "enc:abc"},
        notes="Attention aux doublons.",
    )


def _all_text(doc: Document) -> str:
    chunks = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                chunks.append(cell.text)
    return "\n".join(chunks)


def test_render_job_docx_without_screenshot_falls_back_gracefully(tmp_path: Path):
    job = make_job()
    out_path = tmp_path / "demo-job.docx"

    render_job_docx(job, None, out_path)

    assert out_path.exists()
    doc = Document(str(out_path))
    text = _all_text(doc)
    assert "Demo_Job" in text
    assert "tFileInputDelimited_1" in text
    # Aucune trace du secret en clair, quel que soit le rendu du diagramme.
    assert "enc:abc" not in text
    assert "masqué" in text


def test_render_job_docx_show_context_values(tmp_path: Path):
    job = make_job()
    out_path = tmp_path / "demo-job.docx"

    render_job_docx(job, None, out_path, mask_context_values=False)

    doc = Document(str(out_path))
    assert "enc:abc" in _all_text(doc)


def test_render_job_docx_no_leftover_temp_diagram_file(tmp_path: Path):
    job = make_job()
    out_path = tmp_path / "demo-job.docx"

    render_job_docx(job, None, out_path)

    leftovers = list(tmp_path.glob("*.tmp-diagram.png"))
    assert leftovers == []
