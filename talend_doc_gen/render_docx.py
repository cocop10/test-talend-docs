"""Rendu Word (.docx) d'un job Talend.

Pensé pour être déposé tel quel dans une bibliothèque SharePoint : mise en
page soignée (fiche d'identité, tableaux colorés par catégorie de
composant, callout pour les notes), et diagramme intégré en image — Word
ne sachant pas afficher du Mermaid, on réutilise soit la capture d'écran
déjà résolue par le CLI, soit un rendu du diagramme Mermaid en PNG (voir
:mod:`render_diagram_image`), avec un repli textuel si aucun des deux n'est
disponible.
"""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from .component_style import CATEGORY_STYLE, component_category
from .diagram import build_mermaid
from .models import Job
from .render_diagram_image import render_mermaid_to_png

_ACCENT = RGBColor(0x1D, 0x4E, 0xD8)
_TEXT = RGBColor(0x1B, 0x1F, 0x24)
_MUTED = RGBColor(0x6B, 0x72, 0x80)
_WHITE = RGBColor(0xFF, 0xFF, 0xFF)
_HEADER_BG = "1D4ED8"
_LABEL_BG = "EFF3FA"
_ZEBRA_BG = "F1F5F9"
_BORDER = "D9DEE4"
_NOTE_BG = "FEF9C3"
_NOTE_BORDER = "CA8A04"
_MASKED_VALUE = "•••• (masqué)"

_KIND_STYLE = {
    "flow": (RGBColor(0x16, 0xA3, 0x4A), "Flux principal"),
    "reject": (RGBColor(0xDC, 0x26, 0x26), "Rejet"),
    "trigger": (RGBColor(0x64, 0x74, 0x8B), "Déclencheur"),
    "iterate": (RGBColor(0x25, 0x63, 0xEB), "Itération"),
}


def _hex(color: str) -> str:
    return color.lstrip("#")


def _set_cell_background(cell, hex_color: str) -> None:
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), _hex(hex_color))
    cell._tc.get_or_add_tcPr().append(shd)


def _set_cell_borders(cell, color: str = _BORDER, sz: int = 4) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), str(sz))
        el.set(qn("w:color"), _hex(color))
        borders.append(el)
    tc_pr.append(borders)


def _cell_text(cell, text: str, *, bold=False, italic=False, color=None, size=10) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color


def _heading(doc, text: str, level: int = 2):
    heading = doc.add_heading(level=level)
    run = heading.add_run(text)
    run.font.color.rgb = _ACCENT if level <= 1 else _TEXT
    return heading


def _muted_paragraph(doc, text: str):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(9.5)
    run.font.color.rgb = _MUTED
    return p


def _header_row(table, headers: list[str]) -> None:
    for cell, text in zip(table.rows[0].cells, headers):
        _cell_text(cell, text, bold=True, color=_WHITE)
        _set_cell_background(cell, _HEADER_BG)


def _add_callout(doc, text: str) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    _set_cell_background(cell, _NOTE_BG)
    _set_cell_borders(cell, color=_NOTE_BORDER, sz=6)
    _cell_text(cell, text, size=10)


def _add_framed_image(doc, image_path: Path, width_cm: float = 16.5) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    _set_cell_borders(cell)
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.add_run().add_picture(str(image_path), width=Cm(width_cm))


def _fill_system_cell(cell, title: str, items: list[str]) -> None:
    cell.text = ""
    title_run = cell.paragraphs[0].add_run(title)
    title_run.bold = True
    if items:
        for item in items:
            cell.add_paragraph(item, style="List Bullet")
    else:
        p = cell.add_paragraph()
        run = p.add_run("Non renseigné")
        run.italic = True
        run.font.color.rgb = _MUTED


def _add_category_legend(doc, present_categories: set[str]) -> None:
    ordered = [c for c in CATEGORY_STYLE if c in present_categories]
    if not ordered:
        return
    p = doc.add_paragraph()
    for i, category in enumerate(ordered):
        style = CATEGORY_STYLE[category]
        if i:
            p.add_run("    ")
        swatch = p.add_run("■ ")
        swatch.font.size = Pt(9)
        swatch.font.color.rgb = RGBColor.from_string(_hex(style["stroke"]))
        label = p.add_run(style["label_fr"])
        label.font.size = Pt(9)
        label.font.color.rgb = _MUTED


def render_job_docx(
    job: Job,
    screenshot_path: Path | None,
    out_path: Path,
    mask_context_values: bool = True,
) -> None:
    doc = Document()

    section = doc.sections[0]
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = _TEXT

    footer_p = section.footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_p.add_run("Documentation générée automatiquement — talend-doc-gen")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = _MUTED

    title = doc.add_heading(level=0)
    title_run = title.add_run(job.name)
    title_run.font.color.rgb = _ACCENT

    subtitle_bits = [f"Source du parsing : {job.source_kind}"]
    if job.version:
        subtitle_bits.append(f"Version {job.version}")
    _muted_paragraph(doc, " · ".join(subtitle_bits))

    if job.description:
        p = doc.add_paragraph()
        run = p.add_run(job.description)
        run.font.size = Pt(11)

    _heading(doc, "Fiche d'identité")
    meta_table = doc.add_table(rows=0, cols=2)
    meta_table.style = "Table Grid"
    for label, value in (
        ("Propriétaire", job.owner or "Non renseigné"),
        ("Version", job.version or "Non renseigné"),
        ("Planification", job.schedule or "Non renseigné"),
        ("Source du parsing", job.source_kind),
    ):
        row = meta_table.add_row()
        _cell_text(row.cells[0], label, bold=True)
        _set_cell_background(row.cells[0], _LABEL_BG)
        _cell_text(row.cells[1], value)

    if job.purpose:
        _heading(doc, "Objectif métier")
        doc.add_paragraph(job.purpose)

    if job.source_systems or job.target_systems:
        _heading(doc, "Systèmes impliqués")
        systems_table = doc.add_table(rows=1, cols=2)
        left, right = systems_table.rows[0].cells
        _fill_system_cell(left, "Systèmes source", job.source_systems)
        _fill_system_cell(right, "Systèmes cible", job.target_systems)

    _heading(doc, "Schéma du flux")
    image_path = screenshot_path
    caption = "Capture d'écran du canevas Talend." if image_path else None
    tmp_diagram_path: Path | None = None
    if image_path is None:
        tmp_diagram_path = out_path.with_name(out_path.stem + ".tmp-diagram.png")
        if render_mermaid_to_png(build_mermaid(job), tmp_diagram_path):
            image_path = tmp_diagram_path
            caption = "Diagramme généré automatiquement à partir du flux Talend."

    if image_path is not None:
        _add_framed_image(doc, image_path)
        if caption:
            cap = doc.add_paragraph()
            cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = cap.add_run(caption)
            run.italic = True
            run.font.size = Pt(9)
            run.font.color.rgb = _MUTED
    else:
        _muted_paragraph(
            doc,
            "Aucune capture d'écran ni rendu de diagramme disponible pour ce job "
            "(le rendu du diagramme Mermaid nécessite le paquet optionnel "
            "playwright avec un navigateur Chromium installé).",
        )

    if tmp_diagram_path is not None and tmp_diagram_path.exists():
        tmp_diagram_path.unlink()

    if job.components:
        _heading(doc, "Composants")
        present = {component_category(c.type) for c in job.components}
        _add_category_legend(doc, present)
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        _header_row(table, ["Nom", "Type", "Description"])
        for i, comp in enumerate(job.components):
            row = table.add_row()
            _cell_text(row.cells[0], comp.name)
            _cell_text(row.cells[1], comp.type or "?")
            _set_cell_background(row.cells[1], CATEGORY_STYLE[component_category(comp.type)]["fill"])
            _cell_text(row.cells[2], comp.description or "Non renseigné")
            if i % 2 == 1:
                _set_cell_background(row.cells[0], _ZEBRA_BG)
                _set_cell_background(row.cells[2], _ZEBRA_BG)

    if job.connections:
        _heading(doc, "Connexions")
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        _header_row(table, ["De", "Vers", "Type", "Nature"])
        for i, conn in enumerate(job.connections):
            row = table.add_row()
            _cell_text(row.cells[0], conn.source)
            _cell_text(row.cells[1], conn.target)
            _cell_text(row.cells[2], conn.label)
            color, label_fr = _KIND_STYLE.get(conn.kind, (_MUTED, conn.kind))
            cell = row.cells[3]
            cell.text = ""
            p = cell.paragraphs[0]
            dot_run = p.add_run("● ")
            dot_run.font.color.rgb = color
            p.add_run(label_fr)
            if i % 2 == 1:
                for c in list(row.cells[:3]):
                    _set_cell_background(c, _ZEBRA_BG)

    if job.context_params:
        _heading(doc, "Paramètres de contexte")
        if mask_context_values:
            _muted_paragraph(
                doc,
                "Les valeurs (hôtes, identifiants, chemins...) sont masquées par "
                "défaut car potentiellement sensibles ; seuls les noms de "
                "paramètres sont listés.",
            )
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        _header_row(table, ["Nom", "Valeur"])
        for i, (key, value) in enumerate(job.context_params.items()):
            row = table.add_row()
            _cell_text(row.cells[0], key)
            _cell_text(row.cells[1], _MASKED_VALUE if mask_context_values else value)
            if i % 2 == 1:
                for c in list(row.cells):
                    _set_cell_background(c, _ZEBRA_BG)

    if job.notes:
        _heading(doc, "Notes")
        _add_callout(doc, job.notes)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
